"""Build docs/pitch/overwatering.pptx and docs/pitch/script.md from one source.

    uv run --with python-pptx --with pillow --with "qrcode[pil]" python docs/pitch/tools/build_deck.py

Art PNGs come from tools/art.mjs and screenshots from tools/shots.mjs. Every slide's speaker
notes are also written, in order and with timings, to script.md. Text stays in real text boxes
(Fredoka headings, Nunito body) so the deck is editable after File > Import slides in Google Slides.
"""

import io
import json
from pathlib import Path

import qrcode
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

PITCH = Path(__file__).resolve().parents[1]
ART = PITCH / "art"
SHOTS = PITCH / "shots"
BUILD = PITCH / "tools" / ".build"
BUILD.mkdir(exist_ok=True)

# Numbers read off the live site by tools/shots.mjs, so slides, script and demo path match the screen.
F = json.loads((SHOTS / "facts.json").read_text())
LOT, BILL, PLAN, WEEK, LIVE = F["lot"], F["bill"], F["plan"], F["week"], F["live"]

HEAD, BODY = "Fredoka", "Nunito"
WATER, SKY, MIST = "1f86d6", "4cc3f0", "dff3fd"
LAWN, SUN, CORAL, INK, SAND = "2c9a4b", "ffc93c", "e5533f", "12324a", "f7dcae"
GOLD = "c27c00"  # readable "sun" for text
MUTED = "4f6b80"
BG = "eef8fe"
WHITE = "ffffff"

W, H = 13.333, 7.5


def rgb(h):
    return RGBColor.from_string(h)


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]
SCRIPT = []  # (slide number, title, timing, notes)


# ---------- helpers ----------

def slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(bg)
    return s


def text(s, x, y, w, h, paras, size=20, font=BODY, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, line=None, after=0):
    """paras: list of paragraphs; each is a str or a list of (text, {overrides}) runs."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(0.02))
    for i, para in enumerate(paras if isinstance(paras, list) else [paras]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line:
            p.line_spacing = line
        p.space_after = Pt(after)
        for run_text, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = run_text
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = rgb(o.get("color", color))
    return tb


def kicker(s, t, x=0.75, y=0.55, w=9, color=WATER):
    return text(s, x, y, w, 0.4, t.upper(), size=14, bold=True, color=color)


def title(s, t, x=0.75, y=0.9, w=11.8, h=1.0, size=40, color=INK):
    return text(s, x, y, w, h, t if isinstance(t, list) else [t], size=size, font=HEAD, bold=True, color=color, line=0.95)


def source(s, t, x=0.75, y=6.95, w=11.8):
    return text(s, x, y, w, 0.4, t, size=9, color=MUTED)


def png(s, path, x, y, w=None, h=None):
    kw = {}
    if w is not None:
        kw["width"] = Inches(w)
    if h is not None:
        kw["height"] = Inches(h)
    return s.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)


def box(s, x, y, w, h, fill=WHITE, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE, dash=False):
    b = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        b.adjustments[0] = radius
    if fill:
        b.fill.solid()
        b.fill.fore_color.rgb = rgb(fill)
    else:
        b.fill.background()
    if line:
        b.line.color.rgb = rgb(line)
        b.line.width = Pt(2)
        if dash:
            from pptx.enum.dml import MSO_LINE
            b.line.dash_style = MSO_LINE.DASH
    else:
        b.line.fill.background()
    b.shadow.inherit = False
    b.text_frame.text = ""
    return b


def shot(s, name, x, y, w=None, h=None, crop=None, tag=""):
    """Screenshots go in as JPEG (they're photos of a page; keeps the deck small). Framed in white.
    crop = (left, top, right, bottom) in source pixels."""
    src = SHOTS / f"{name}.png"
    if not src.exists():
        return None
    out = BUILD / f"{name}{tag}.jpg"
    im = Image.open(src).convert("RGB")
    if crop:
        im = im.crop(crop)
    im.thumbnail((2400, 2400))
    im.save(out, quality=86)
    ar = im.width / im.height
    if w is None:
        w = h * ar
    if h is None:
        h = w / ar
    pad = 0.07
    box(s, x - pad, y - pad, w + 2 * pad, h + 2 * pad, fill=WHITE, line="cfe6f5", radius=0.03)
    return png(s, out, x, y, w, h)


def placeholder(s, x, y, w, h, label, detail=""):
    box(s, x, y, w, h, fill="fff6e0", line=CORAL, dash=True, radius=0.05)
    paras = [[("PLACEHOLDER", {"bold": True, "color": CORAL, "size": 13})], [(label, {"font": HEAD, "bold": True, "size": 24})]]
    if detail:
        paras.append([(detail, {"size": 15, "color": MUTED})])
    text(s, x + 0.3, y + 0.3, w - 0.6, h - 0.6, paras, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, after=6)


def notes(s, timing, body):
    """Speaker notes: a timing line, then the words. Also collected for script.md."""
    body = "\n".join(line.strip() for line in body.strip().splitlines())
    s.notes_slide.notes_text_frame.text = f"[{timing}]\n{body}"
    SCRIPT.append((len(prs.slides), slide_titles.get(len(prs.slides), ""), timing, body))


slide_titles = {}


def named(t):
    slide_titles[len(prs.slides)] = t


def logo_word(s, x, y, size=40):
    return text(s, x, y, 6, size / 50, [[("Overwater", {}), (".ing", {"color": WATER}), ("?", {"color": "f0a92a"})]],
                size=size, font=HEAD, bold=True)


def big_stat(s, x, y, w, number, label, color=WATER, nsize=88, lsize=20, lh=1.4):
    text(s, x, y, w, nsize / 60, number, size=nsize, font=HEAD, bold=True, color=color, line=0.9)
    text(s, x, y + nsize / 60 + 0.05, w, lh, label, size=lsize, color=INK, line=1.1)


def footer_scene(s):
    png(s, ART / "scene-footer.png", 0, H - W * 440 / 2400, w=W)


def icon(kind, color):
    """A check (done) or an open ring (to do), as PNG so it imports anywhere."""
    from PIL import ImageDraw
    out = BUILD / f"icon-{kind}-{color}.png"
    S = 400
    im = Image.new("RGBA", (S, S), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    c = tuple(int(color[i:i + 2], 16) for i in (0, 2, 4))
    if kind == "check":
        d.ellipse((8, 8, S - 8, S - 8), fill=c)
        d.line([(110, 210), (175, 275), (295, 135)], fill="white", width=44, joint="curve")
    else:
        d.ellipse((20, 20, S - 20, S - 20), outline=c, width=34)
    im.resize((100, 100), Image.LANCZOS).save(out)
    return out


def freeform(s, pts, fill):
    """A filled polygon from (x, y) points in inches."""
    fb = s.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
    fb.add_line_segments([(Inches(x), Inches(y)) for x, y in pts[1:]], close=True)
    shp = fb.convert_to_shape()
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def n(v):
    return f"{v:,}"


def thou(v):
    return f"{round(v, -3):,}"


def qr(url):
    q = qrcode.QRCode(border=1, box_size=20, error_correction=qrcode.constants.ERROR_CORRECT_M)
    q.add_data(url)
    img = q.make_image(fill_color="#12324a", back_color="white")
    out = BUILD / "qr.png"
    img.save(out)
    return out


# ---------- the deck ----------

# 1. Title
s = slide()
named("Title")
png(s, ART / "scene-hero.png", 6.9, 0.9, w=6.2)
png(s, ART / "logo.png", 0.75, 1.15, w=1.3)
logo_word(s, 2.15, 1.42, size=46)
text(s, 0.75, 2.85, 6.4, 2.0, [
    [("Is your lawn getting more water than it needs?", {})],
], size=34, font=HEAD, bold=True, line=1.0)
text(s, 0.75, 4.55, 6.0, 1.2, "Every lawn in Utah, measured from the air. What it really needs, sent to the sprinkler controller.",
     size=20, color=MUTED, line=1.15)
text(s, 0.75, 6.25, 6.4, 0.5, [[("Adam Fortuna", {"bold": True}), ("  ·  overwater.ing  ·  Nucleus × Startup State", {})]],
     size=16, color=INK)
notes(s, "0:00–0:05", """
Hi, I'm Adam Fortuna. This is Overwater.ing. Question mark.
""")

# 2. The story
s = slide()
named("How I got here")
kicker(s, "How I got here")
title(s, "It started over a pint.")
steps = [
    ("July", "Perspectives Over Pints, Fisher Brewing.", "The Yard Farmer gave the talk."),
    ("What I learned", "Utah pays you to replace your grass.", "Utah Water Savers: $1.50–$3 per square foot."),
    ("November", "Our own yard consult is booked.", "So: how much lawn do I actually have?"),
]
for i, (when, a, b) in enumerate(steps):
    x = 0.75 + i * 4.05
    box(s, x, 2.2, 3.75, 2.85, fill=WHITE, line="cfe6f5", radius=0.1)
    text(s, x + 0.3, 2.4, 3.2, 0.45, when.upper(), size=14, bold=True, color=WATER)
    text(s, x + 0.3, 2.85, 3.2, 1.2, a, size=21, font=HEAD, bold=True, line=1.0)
    text(s, x + 0.3, 4.1, 3.2, 0.85, b, size=17, color=MUTED, line=1.1)
text(s, 0.75, 5.3, 8.0, 1.2, [[("Nobody could tell me. ", {"bold": True}), ("Not for my lot, and not for any lot in Utah.", {})]],
     size=24, line=1.1)
png(s, ART / "grass.png", 10.9, 4.95, w=1.6)
png(s, ART / "drop.png", 10.2, 5.65, w=0.62)
notes(s, "0:05–0:30", """
In July I went to Perspectives Over Pints at Fisher Brewing, and The Yard Farmer gave the talk.
That's where I learned about Utah Water Savers: the state will pay you up to three dollars a
square foot to replace your grass. We booked our own yard consult for November. But my first
question was: how much lawn do I have, and what does it really need? Nobody could tell me.
""")

# 3. Where Utah's water goes (two stacked bars: all water, then cities' water)
s = slide()
named("Where Utah's water goes")
kicker(s, "The problem")
title(s, "Where Utah's water goes.")
X0, BW = 0.75, 11.85


def bar(s, y, h, segs, label, gap=0.04, lw=BW):
    """segs: (name, value, fill, text color or None for no inside label). Returns segment x-ranges."""
    text(s, X0, y - 0.42, lw, 0.4, label, size=15, bold=True, color=INK)
    total = sum(v for _, v, _, _ in segs)
    x, spans = X0, {}
    for name, v, fill, tc in segs:
        w = BW * v / total
        r = box(s, x, y, max(w - gap, 0.02), h, fill=fill, radius=0.0, shape=MSO_SHAPE.RECTANGLE)
        if tc and w > 1.2:
            text(s, x + 0.12, y, w - 0.2, h, [[(name, {"bold": True}), (f"  {100 * v / total:.0f}%", {})]],
                 size=15, color=tc, anchor=MSO_ANCHOR.MIDDLE)
        spans[name] = (x, x + w - gap)
        x += w
    return spans


# USGS Circular 1441, Table 2B (Utah, 2015, thousand acre-feet a year, fresh + saline)
ALL = [
    ("Farms", 3390, "e6c590", INK),
    ("Cities & towns", 714, WATER, WHITE),
    ("Mining", 293, "9fb4c4", None),
    ("Industry", 149, "b3c4d1", None),
    ("Fish hatcheries", 93, "c6d3dd", None),
    ("Power plants", 78, "d6e0e8", None),
    ("Livestock", 18, "e3eaf0", None),
]
spans = bar(s, 2.35, 0.8, ALL, "All of Utah's water: ~4.7M acre-feet a year", lw=5.6)
text(s, X0 + BW - 6.1, 1.97, 6.1, 0.35, "Then mining 6%, industry 3%, fish hatcheries 2%, power 2%, livestock 0.4%",
     size=12, color=MUTED, align=PP_ALIGN.RIGHT)
cx0, cx1 = spans["Cities & towns"]
# zoom wedge from the cities segment to the second bar
freeform(s, [(cx0, 3.15), (cx1, 3.15), (X0 + BW, 4.25), (X0, 4.25)], "d4ecfb")
# Utah DWR 2021 (via USU Extension): residential is >70% of M&I, 60% of that outdoors
CITY = [
    ("Home yards", 42, LAWN, WHITE),
    ("Inside homes", 28, "a9d6f5", INK),
    ("Businesses, schools, churches & industry", 30, "cddbe5", INK),
]
bar(s, 4.7, 0.8, CITY, "Where cities' water goes, counting secondary (untreated) irrigation water")
text(s, X0, 5.62, BW, 1.1, [
    [("Home yards: 500,993 acre-feet a year,", {"bold": True, "color": LAWN})],
    [("enough to raise Great Salt Lake about 10 inches.", {"bold": True})],
], size=22, font=HEAD, line=1.0)
source(s, "All water: USGS Circular 1441, Utah 2015 (withdrawals incl. saline; much of Utah's secondary yard water is counted under irrigation there). Cities: Utah DWR Water Resources Plan 2021 via USU Extension (M&I is 25% of 4.75M AF diverted; >70% residential; 60% of that outdoors). Lake: ~610,000 acres near 4,191 ft (UGS), volume equivalent only; saved water can increase potential flows to the lake.", y=6.8)
notes(s, "0:30–0:45", """
Here's all of Utah's water. Most goes to farms. About a sixth runs through our cities, and the
biggest share of that is home yards: half a million acre-feet a year. Enough to raise Great Salt
Lake about ten inches.
""")

# 4. The lake needs it back
s = slide()
named("The lake needs that water back")
kicker(s, "Why now")
title(s, "The lake needs that water back.")
box(s, 0.75, 2.2, 5.7, 3.75, fill=WHITE, line="cfe6f5", radius=0.06)
text(s, 1.05, 2.4, 5.2, 0.4, "GREAT SALT LAKE", size=14, bold=True, color=WATER)
text(s, 1.05, 2.8, 5.2, 1.2, "4,190.4 ft", size=64, font=HEAD, bold=True, color=CORAL, line=0.9)
text(s, 1.05, 4.0, 5.2, 0.6, "August 2026. The healthy range starts at 4,198 ft.", size=18, line=1.1)
text(s, 1.05, 4.75, 5.2, 1.2, [[("471,000–1,055,000 ", {"bold": True}), ("more acre-feet a year to get it back.", {})]],
     size=18, line=1.1)
box(s, 6.8, 2.2, 5.8, 3.75, fill=WHITE, line="cfe6f5", radius=0.06)
text(s, 7.1, 2.4, 5.2, 0.4, "AND IT STAYS WITH THE LAKE", size=14, bold=True, color=LAWN)
text(s, 7.1, 2.85, 5.2, 1.0, "Water saved on yards can flow to the lake.", size=26, font=HEAD, bold=True, line=1.0)
text(s, 7.1, 4.0, 5.25, 1.9, [
    [("Executive Order 2026-03 ", {"bold": True}), ("(May 2026): new data centers may not add to Great Salt Lake's depletion.", {})],
    [("So Utah's AI build-out can't take what yards give back.", {"bold": True})],
], size=18, line=1.1, after=10)
png(s, ART / "arch-solo.png", 11.55, 0.55, h=1.4)
source(s, "Sources: Utah News Dispatch / USGS gauge (lake level); Great Salt Lake Strategic Plan (inflow need); Utah Executive Order 2026-03 (May 29, 2026). Saved water can increase potential flows to the lake; it may be reused upstream.")
notes(s, "0:45–1:10", """
Great Salt Lake is eight feet below healthy. Water we save on yards can flow to the lake. And
since May, new data centers can't add to its depletion, so the AI build-out can't take it.
But nobody knows, lot by lot, how much lawn there is or what it really needs. So I built it.
""")

# 5. Demo: the homeowner's flow
s = slide()
named("Live demo")
kicker(s, "Live demo · overwater.ing")
title(s, "Type any Utah address.")
COLS = [0.75, 4.95, 9.15]
CW = 3.45
steps = [
    ("1", "Enter your address", f"Any address in all 29 counties."),
    ("2", "Confirm what's grass", f"Tap a shape to fix it. {n(LOT['lawn_sqft'])} sq ft of lawn here."),
    ("3", "See what native saves", f"{n(LOT['native_gal'])} gallons a year, and up to ${n(LOT['rebate_usd'])} back."),
]
for (num, head, cap), x in zip(steps, COLS):
    text(s, x, 1.95, CW, 0.5, [[(num + "  ", {"font": HEAD, "color": WATER}), (head, {})]], size=21, font=HEAD, bold=True)
    text(s, x, 4.95, CW, 0.5, cap, size=14, color=MUTED, line=1.05)
for x in COLS[:2]:
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + CW + 0.22), Inches(3.45), Inches(0.42), Inches(0.32))
    a.fill.solid(); a.fill.fore_color.rgb = rgb(SKY); a.line.fill.background(); a.shadow.inherit = False
shot(s, "flow-1-search", COLS[0] + 0.05, 3.15, w=CW - 0.1)
shot(s, "flow-2-map", COLS[1] + (CW - 2.35) / 2, 2.55, h=2.3)
shot(s, "flow-3-native", COLS[2] + 0.05, 2.75, w=CW - 0.1)
box(s, 0.75, 5.6, 11.85, 1.45, fill=MIST, radius=0.14)
text(s, 1.05, 5.7, 9.2, 1.3, [
    [("4  ", {"color": WATER}), ("Or lower your water use today", {})],
    [("County-level watering schedules that follow the state's weekly guide, as a calendar or sent straight to the controller: OpenSprinkler, Rachio or Home Assistant.",
      {"font": BODY, "bold": False, "size": 15, "color": INK})],
], size=21, font=HEAD, bold=True, line=1.1, after=4)
png(s, ART / "sprinkler.png", 10.7, 5.55, h=1.4)
notes(s, "1:10–3:10", f"""
(Switch to the browser. Full click path and fallbacks are in demo.md.)
1:10 Type "603 E South Temple", the Governor's Mansion. Pick it.
1:20 The 2024 aerial photo. Green is lawn, dark green is trees, red is buildings: {n(LOT['lawn_sqft'])} square
feet of lawn. That outline comes from 60-centimeter infrared aerial photos, minus the trees from
laser-scanned tree heights, minus the roofs, computed right in your browser. If we got a shape
wrong, tap it, and that correction trains the model. (Tap one, then tap it back to green.)
1:40 Go native: about {thou(LOT['native_gal'])} gallons a year saved. And Utah pays up to ${n(LOT['rebate_usd'])} here,
from {LOT['rebate_program']}. It knows the city, the water system and the program.
1:55 What going native looks like: the before and after, the plants, the designers to call.
2:05 Or lower your water use today: pick spray heads and a basic timer. One thing to change, and a
month-by-month schedule for this county.
2:20 Subscribe to it as a calendar, or send it to the controller. This is OpenSprinkler's real
firmware: when the state says skip, it switches itself off. In July: Monday, Wednesday, Friday at
5 a.m., four short cycles a zone.
2:50 (Only if the warm-up worked.) Any address in Utah is measured live, like this one in St. George.
""")

# 6. Three phases: what $10K unlocks
s = slide()
named("What $10K unlocks")
kicker(s, "The plan")
title(s, "Three phases. $10K unlocks the second.")
phases = [
    ("PHASE 1 · LIVE NOW", LAWN, "check", [
        "Estimate your rebate to go native",
        "Optimize your watering schedule",
        "OpenSprinkler integration, plus Rachio and Home Assistant",
    ], None),
    ("PHASE 2 · WHAT $10K FUNDS", WATER, "ring", [
        "Pilot: manage sprinklers remotely",
        "Support more sprinkler systems",
        "Moisture sensor integration",
        "Meet Utah's water-wise partners: what's in progress, next steps",
    ], "slowtheflow.org/partners"),
    ("PHASE 3 · WITH THE STATE", GOLD, "ring", [
        "Drone scans tune each sprinkler head",
    ], "State drone flights a few times a week map moisture with LiDAR. Each sprinkler waters only the patch that needs it, and skips the rest."),
]
for i, (head, c, kind, items, foot) in enumerate(phases):
    x = 0.75 + i * 4.05
    box(s, x, 1.95, 3.75, 4.95, fill=WHITE, line=c if i == 1 else "cfe6f5", radius=0.06)
    text(s, x + 0.25, 2.12, 3.3, 0.4, head, size=13, bold=True, color=c)
    y = 2.65
    for it in items:
        png(s, icon(kind, LAWN if kind == "check" else c), x + 0.25, y + 0.04, w=0.32)
        lines = 1 + len(it) // 25
        text(s, x + 0.7, y, 2.85, 0.33 * lines + 0.1, it, size=15, bold=True, line=1.05)
        y += 0.33 * lines + 0.26
    if foot:
        text(s, x + 0.7 if i == 1 else x + 0.25, y - 0.05 if i == 1 else y + 0.1, 3.0, 1.6, foot, size=12 if i == 1 else 15, color=MUTED, line=1.15)
notes(s, "3:10–3:45", """
Phase one is live today: your rebate estimate, a watering schedule, and connectors for
OpenSprinkler, Rachio and Home Assistant. Ten
thousand dollars funds phase two: a pilot that manages real sprinklers remotely, more sprinkler
systems, moisture sensors, and sitting down with each of Utah's water-wise partners to see what's
already in progress. Phase three is with the state: drone scans a few times a week, so each
sprinkler head waters only the patch that needs it, and the state can spot overwatering from
the air before a bill ever shows it.
""")

# 7. Ask + close
s = slide()
named("The ask")
footer_scene(s)
text(s, 0.75, 0.75, 8.0, 0.4, "THE ASK", size=14, bold=True, color=WATER)
text(s, 0.75, 1.15, 8.2, 1.9, "Type your address at overwater.ing", size=48, font=HEAD, bold=True, line=0.95)
text(s, 0.75, 3.05, 8.4, 1.9, [
    [("$10K ", {"bold": True, "color": WATER}), ("for phase 2: a pilot that runs real sprinklers.", {})],
    [("Intros ", {"bold": True, "color": WATER}), ("to Utah's water-wise partners, a district conservation team, and drone training academies for LiDAR scans.", {})],
], size=20, line=1.15, after=8)
png(s, qr("https://overwater.ing"), 9.75, 0.85, w=2.6)
text(s, 9.75, 3.5, 2.6, 0.4, "overwater.ing", size=18, font=HEAD, bold=True, align=PP_ALIGN.CENTER)
notes(s, "3:45–4:05", """
So here's the ask: ten thousand dollars for phase two, a pilot that runs real sprinklers. And
introductions: to Utah's water-wise partners, to a district's conservation team, and to drone
training academies that could fly LiDAR scans with us. And right now, take out your phone and
type your address at overwater.ing. Thank you.
""")

# ---------- backup / appendix ----------

s = slide(MIST)
named("Backup")
text(s, 0.75, 2.6, 11.8, 1.2, "Backup: the demo, step by step", size=44, font=HEAD, bold=True)
text(s, 0.75, 3.75, 11.8, 0.8, "If the Wi-Fi fails, click through these instead. Same order as the live demo.", size=20, color=MUTED)
png(s, ART / "sprinkler-off.png", 10.8, 4.6, w=1.4)
notes(s, "backup", "Only if the live demo fails. Each following slide matches one demo step in demo.md.")


def backup(name, head, caption, timing, note, zoom=None):
    """The page on the left; a close-up of the part that matters and the caption on the right."""
    s = slide()
    named(head)
    text(s, 0.75, 0.4, 11.8, 0.7, head, size=30, font=HEAD, bold=True)
    sh = 5.15
    if not zoom:
        shot(s, name, (W - sh * 1.6) / 2, 1.3, h=sh)
        if caption:
            text(s, 0.75, 6.65, 11.8, 0.6, caption, size=15, color=MUTED, align=PP_ALIGN.CENTER)
    else:
        shot(s, name, 0.75, 1.3, h=sh)
        zx, zw = 9.4, 3.2
        zname, zcrop = zoom if isinstance(zoom, tuple) else (zoom, None)
        im = Image.open(SHOTS / f"{zname}.png")
        if zcrop:
            im = im.crop(zcrop)
        zh = min(zw * im.height / im.width, 3.7)
        zw2 = zh * im.width / im.height
        shot(s, zname, zx + (zw - zw2) / 2, 1.3, w=zw2, crop=zcrop, tag="-zoom")
        if caption:
            text(s, zx, 1.3 + zh + 0.3, zw, 6.45 - (1.3 + zh + 0.3), caption, size=14, color=INK, line=1.15)
    notes(s, timing, note)
    return s


backup("02-search", "1 · Enter any Utah address", "Search covers parcels in all 29 counties.", "demo 1:10",
       "Type 603 E South Temple and pick the first result.", zoom="flow-1-search")
backup("03-lot", "2 · Confirm what's grass",
       f"2024 USDA imagery at 0.6 m. Green = lawn, dark green = trees, red = buildings. Tap a shape to fix it. {n(LOT['lawn_sqft'])} sq ft of lawn here.",
       "demo 1:20", f"{n(LOT['lawn_sqft'])} sq ft of lawn on the Governor's Mansion lot, about {LOT['need_in']} inches and {thou(LOT['need_gal'])} gallons a year.", zoom="flow-2-lawn")
backup("07-native", "3 · What going native saves",
       f"About {n(LOT['native_gal'])} gallons a year for this lot, then the plants and the designers to call.", "demo 1:55",
       "The before and after for this lot.", zoom="flow-3-native")
backup("08-rebate", "3 · The rebate, and who's in charge here",
       f"Up to ${n(LOT['rebate_usd'])} from {LOT['rebate_program']}. City, water provider and program for this exact lot.", "demo 1:40",
       "The city, the water system and the program for this lot, plus the $100 smart controller rebate.", zoom="flow-3-rebate")
backup("05-watering-plan", "4 · Lower your water use today", "Spray heads and a basic timer, in this county's weather: one thing to change.", "demo 2:05",
       "The single change with the most gallons saved, then the month-by-month schedule.", zoom=("flow-4-plan", (68, 45, 1410, 300)))
backup("05-watering-schedule", "4 · The schedule", "Days a week and cycle-and-soak minutes per zone, from local weather.", "demo 2:10",
       "Deep and infrequent, split into short cycles so it soaks in.")
backup("06-calendar", "4 · Subscribe to the watering calendar", "It follows Utah's Weekly Lawn Watering Guide; skip weeks clear themselves.", "demo 2:20",
       "Subscribe once; each watering day shows up at 5 a.m. with run times.")

# 4 · the controller, off and on (OpenSprinkler firmware in demo mode, set by our connector)
s = slide()
named("4 · The controller follows the schedule")
text(s, 0.75, 0.4, 11.8, 0.7, "4 · The controller follows the schedule", size=30, font=HEAD, bold=True)
text(s, 0.75, 1.25, 5.7, 0.6, "When the state says skip: off", size=26, font=HEAD, bold=True, color=CORAL)
text(s, 6.9, 1.25, 5.7, 0.6, "A July week: on", size=26, font=HEAD, bold=True, color=LAWN)
shot(s, "06-os-paused-full", 0.75, 1.95, w=5.6, crop=(300, 320, 1500, 710), tag="-enabled")
shot(s, "06-os-july-full", 6.9, 1.95, w=5.6, crop=(300, 410, 1500, 800), tag="-enabled")
text(s, 0.75, 4.05, 5.6, 1.5, "Utah's Weekly Lawn Watering Guide said no watering in Salt Lake County, so the connector switched the program off. Other programs are untouched.",
     size=17, line=1.15)
shot(s, "06-os-july-full", 6.9, 4.05, w=3.5, crop=(300, 1500, 1500, 2430), tag="-stations")
text(s, 10.6, 4.05, 2.1, 2.7, "Mon, Wed, Fri at 5:00 AM. 4 × 7 min per zone, an hour apart, so it soaks in.", size=15, line=1.15)
text(s, 0.75, 6.2, 5.6, 0.8, "OpenSprinkler's own firmware, in demo mode, set by overwater.ing/connect. One program, rewritten each morning.", size=13, color=MUTED, line=1.1)
notes(s, "demo 2:25", """
OpenSprinkler firmware in demo mode, set by our connector. When the guide says no watering, the program is switched
off. Preview a July week and it's Monday, Wednesday, Friday at 5 a.m., four 7-minute cycles per zone.
""")
backup("06-connect", "4 · Any controller: /connect", "Pick a county and sprinkler type, then OpenSprinkler, Rachio or Home Assistant. The same schedule is open JSON.", "demo 2:30",
       "Rachio and Home Assistant cover most smart controllers sold in Utah; OpenSprinkler is the open-hardware path.")
backup("04-bill", f"Optional · Claude reads the bill: {n(BILL['gap_gal'])} gallons over",
       f"{n(BILL['outdoor_gal'])} gallons a year outdoors against the {n(LOT['need_gal'])} this lawn needs, read from a photo of the bill.", "optional",
       "Upload the sample bill. Claude reads twelve months; the page shows the overwatering gap.", zoom="zoom-bill")
backup("09-live", "Optional · Any address, measured live",
       f"67 W 200 N, St. George: {n(LIVE['lawn_sqft'])} sq ft of lawn needing {LIVE['need_in']} inches a season. Measured in the browser, nothing precomputed.", "demo 2:50",
       "Any Utah address is measured live in the browser in seconds.", zoom="zoom-live")
backup("11-district-avenues", "Neighborhood view for districts", "Totals only, never household rankings.", "backup",
       "What a district sees: lawn, need, likely overwatering and savings for a whole neighborhood.")
backup("12-impact", "Utah-scale numbers: overwater.ing/impact", "Our measured neighborhoods and the state-scale scenario.", "backup",
       "Utah-scale numbers, live on the site.")

# Sources
s = slide()
named("Sources")
text(s, 0.75, 0.4, 11.8, 0.7, "Sources", size=30, font=HEAD, bold=True)
SOURCES = [
    "All Utah water by category (2015): USGS Circular 1441, Estimated Use of Water in the United States in 2015, Table 2B",
    "Cities' water: Utah DWR Water Resources Plan (Dec 2021) via USU Extension, extension.usu.edu/climate/research/water-wise-policies-and-programs-utah: 4.75M AF diverted, M&I 25%, >70% residential, 60% outdoors, 500,993 AF/yr",
    "Great Salt Lake area ~610,000 acres near 4,191 ft: Utah Geological Survey. Level 4,190.4 ft (Aug 2026), healthy from 4,198 ft: Utah News Dispatch; wre-rapps.utah.gov/GSLApp",
    "471,000–1,055,000 AF/yr more inflow: Great Salt Lake Strategic Plan, greatsaltlake.utah.gov",
    "Data centers may not add GSL depletion: Utah Executive Order 2026-03 (May 29, 2026)",
    "Rebates of $1.50–$3 per sq ft: Utah Water Savers Landscape Incentive Program (rates checked Oct 27 2026)",
    "Water-wise partners: slowtheflow.org/partners",
]
text(s, 0.75, 1.25, 11.8, 5.9, [[("•  " + t, {})] for t in SOURCES], size=13, line=1.1, after=5)
notes(s, "reference", "Sources for every number in the deck. Saved water can increase potential flows to the lake; it may be reused upstream.")

out = PITCH / "overwatering.pptx"
prs.save(out)
print("wrote", out, f"{out.stat().st_size / 1e6:.1f} MB,", len(prs.slides), "slides")

# ---------- script.md ----------
md = ["# Pitch script: Overwater.ing?", "",
      "Generated by `docs/pitch/tools/build_deck.py` from the same text as the slide notes, with the demo",
      f"lot's numbers as read from the live site on {LOT['checked'][:10]} (`shots/facts.json`). Target: 5:00. Click path: `demo.md`.", "",
      "| # | Slide | Time |", "|---|---|---|"]
main = [row for row in SCRIPT if row[2][0].isdigit()]
for num, t, timing, _ in main:
    md.append(f"| {num} | {t} | {timing} |")
md.append("")
for num, t, timing, body in main:
    md += [f"## {timing} · Slide {num}: {t}", "", body, ""]
md += ["## Backup slides", "", "Only if the live demo fails. They follow the demo order.", ""]
for num, t, timing, body in SCRIPT:
    if not timing[0].isdigit():
        md.append(f"- **Slide {num}: {t}** ({timing}). {body}")
(PITCH / "script.md").write_text("\n".join(md) + "\n")
print("wrote script.md")

# ---------- demo.md, from tools/demo.template.md ----------
slide_no = {t: num for num, t, _, _ in SCRIPT}
week_off = WEEK["paused"]
ctx = {
    **{f"lot_{k}": (n(v) if isinstance(v, int) else v) for k, v in LOT.items()},
    **{f"live_{k}": (n(v) if isinstance(v, int) else v) for k, v in LIVE.items()},
    "bill_gap": n(BILL["gap_gal"]), "bill_outdoor": n(BILL["outdoor_gal"]),
    "plan_one_thing": PLAN["one_thing"], "plan_save": n(PLAN["save_gal"]),
    "week_text": WEEK["text"],
    "off_cmd": "scripts/opensprinkler-demo.sh sync" if week_off else "scripts/opensprinkler-demo.sh sync --date 2026-12-09",
    "off_says": "This week the state says no watering, so it switched itself off." if week_off else "In the off season, it switches itself off. (The state guide wasn't in the schedule when this was built, so this uses a December week.)",
    "lot_native_k": thou(LOT["native_gal"]), "lot_need_k": thou(LOT["need_gal"]),
    "checked": LOT["checked"][:16].replace("T", " ") + " UTC",
    **{f"s_{k}": v for k, v in {
        "search": slide_no["1 · Enter any Utah address"], "lot": slide_no["2 · Confirm what's grass"],
        "native": slide_no["3 · What going native saves"], "rebate": slide_no["3 · The rebate, and who's in charge here"],
        "plan": slide_no["4 · Lower your water use today"], "cal": slide_no["4 · Subscribe to the watering calendar"],
        "os": slide_no["4 · The controller follows the schedule"], "live": slide_no["Optional · Any address, measured live"],
        "back": slide_no["Backup"], "demo": slide_no["Live demo"], "last": len(prs.slides) - 1, "after_demo": slide_no["What $10K unlocks"],
    }.items()},
}
tpl = (PITCH / "tools" / "demo.template.md").read_text()
(PITCH / "demo.md").write_text(tpl.format_map(ctx))
print("wrote demo.md")
